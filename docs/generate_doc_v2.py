"""
生成智能停车场管理系统项目设计书 Word 文档（精简版——仅列技术方案，不展开实现细节）
"""
import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    shading_elm.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(cell, '2B579A')
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = str(val)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
            if r % 2 == 1:
                set_cell_shading(cell, 'F2F6FC')
    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(width)
    return table

def generate_doc():
    doc = Document()

    # ===== 页面设置 =====
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    # ===== 默认样式 =====
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(4)

    for i in range(1, 4):
        h_style = doc.styles[f'Heading {i}']
        h_style.font.name = '黑体'
        h_style.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        h_style.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)
        if i == 1: h_style.font.size = Pt(20)
        elif i == 2: h_style.font.size = Pt(15)
        else: h_style.font.size = Pt(13)

    # ============ 封面 ============
    for _ in range(7):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('智能停车场管理系统')
    r.font.name = '黑体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    r.font.size = Pt(36); r.font.color.rgb = RGBColor(0x2B, 0x57, 0x9A); r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('项 目 设 计 书')
    r.font.name = '黑体'; r.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    r.font.size = Pt(26); r.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)

    for _ in range(4): doc.add_paragraph()

    for label, val in [('版本号', 'V1.0'), ('编制日期', '2026年6月'), ('文档密级', '内部')]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f'{label}：{val}'); r.font.size = Pt(14)

    doc.add_page_break()

    # ============ 目录 ============
    doc.add_heading('目  录', level=1)
    for item, pg in [('一、项目概述', 3), ('二、系统总体架构', 4), ('三、车牌识别模块（边缘端）', 5),
                      ('四、云端识别服务模块', 7), ('五、后端服务模块', 8), ('六、数据库模块', 9),
                      ('七、AI 预测与推荐模块', 10), ('八、前端模块', 11), ('九、硬件通信模块', 12),
                      ('十、计费引擎模块', 13), ('十一、实时推送模块', 14), ('十二、部署与运维模块', 15),
                      ('十三、测试方案', 16), ('十四、项目里程碑', 17)]:
        p = doc.add_paragraph()
        r = p.add_run(f'{item}{"." * (55 - len(item))}{pg}'); r.font.size = Pt(12)

    doc.add_page_break()

    # ================================================================
    # 一、项目概述
    # ================================================================
    doc.add_heading('一、项目概述', level=1)
    doc.add_paragraph(
        '本项目构建一套集车牌识别、智能推荐、实时监控、自动计费于一体的'
        '智能停车场管理系统，实现停车场无人化、智能化运营。系统采用"边缘计算 + 云端服务 + Web 前端"三层架构。'
    )

    # ================================================================
    # 二、系统总体架构
    # ================================================================
    doc.add_heading('二、系统总体架构', level=1)

    doc.add_heading('2.1 技术栈总览', level=2)
    add_table(doc,
        ['层级', '技术选型'],
        [
            ['边缘识别硬件', 'K210 (AI视觉) + ESP32-S3 (通信控制)'],
            ['边缘识别算法', 'YOLO (车辆检测) + 传统 CV (车牌定位) + CNN (字符分类 OCR)'],
            ['云端 OCR 兜底', 'PaddleOCR / 腾讯云 LicensePlateOCR API'],
            ['后端框架', 'FastAPI (推荐) / Flask + SQLAlchemy'],
            ['数据库', 'PostgreSQL (持久化) + Redis (缓存/实时)'],
            ['AI 预测', 'LSTM / Prophet (车流时序预测)'],
            ['AI 推荐', '协同过滤 Collaborative Filtering (车位推荐)'],
            ['前端框架', 'Vue 3 + Vite + TypeScript + Element Plus + ECharts'],
            ['实时通信', 'WebSocket (flask-socketio / socket.io)'],
            ['硬件通信', 'PySerial (本地串口) 或 MQTT (远程)'],
            ['部署', 'Docker + Gunicorn + Nginx + Let\'s Encrypt'],
            ['测试', 'pytest + Playwright'],
            ['监控', 'Prometheus + Grafana + Sentry'],
        ],
        col_widths=[4, 13]
    )

    doc.add_heading('2.2 数据流概述', level=2)
    doc.add_paragraph(
        '入场：摄像头 → K210 (YOLO 检测车辆 → CV 定位车牌 → CNN 识别字符) → ESP32 (WiFi) → '
        '云端 API → 模糊匹配注册车牌 → 分配车位 → 开闸 → 写入停车记录。'
    )
    doc.add_paragraph(
        '出场：摄像头 → 识别车牌 → 查找在停记录 → 计费引擎计算费用 → 支付 → 开闸 → 释放车位。'
    )

    # ================================================================
    # 三、车牌识别模块（边缘端）
    # ================================================================
    doc.add_heading('三、车牌识别模块（边缘端）', level=1)

    doc.add_heading('3.1 硬件方案', level=2)
    add_table(doc,
        ['硬件', '型号', '职责'],
        [
            ['AI 视觉芯片', 'K210 (Kendryte)', '运行 YOLO 车辆检测 + CNN 字符识别 (KPU 硬件加速)'],
            ['WiFi 通信 MCU', 'ESP32-S3', '数据上传、闸机控制、OLED 显示、时间同步'],
            ['摄像头模组', 'OV2640 / OV5640', '图像采集'],
            ['补光灯', 'LED 白光 + 红外', '夜间补光（光敏电阻自动触发）'],
        ],
        col_widths=[3, 4, 10]
    )

    doc.add_heading('3.2 算法流水线', level=2)
    add_table(doc,
        ['步骤', '技术方案', '说明'],
        [
            ['车辆检测', 'YOLOv2-tiny 目标检测', '判断画面中是否出现车辆，过滤行人/非机动车干扰'],
            ['车牌定位', '传统 CV（颜色分割 + 轮廓检测 + 透视变换）', '从车辆画面中提取车牌区域并校正倾斜'],
            ['字符识别', 'CNN 分类器（LeNet-5 改进版，KPU 推理）', '逐字符识别，输出车牌号 + 置信度'],
            ['协同策略', '边缘置信度 ≥ 0.85 直接使用；< 0.85 上传原图到云端兜底', '本地优先、云端备用'],
            ['去抖机制', '滑动窗口：5 帧内 ≥ 3 帧相同结果才确认', '过滤误识别和噪声'],
        ],
        col_widths=[2.5, 6, 8.5]
    )

    doc.add_heading('3.3 边缘-云端协同', level=2)
    add_table(doc,
        ['场景', '处理方式'],
        [
            ['边缘置信度 ≥ 0.85', '直接上传结果，无需云端'],
            ['边缘置信度 < 0.85', '上传原始图像，触发云端 OCR'],
            ['网络断开', '本地缓存结果，恢复后批量补传'],
            ['云端超时', '降级使用边缘结果'],
        ],
        col_widths=[6, 11]
    )

    # ================================================================
    # 四、云端识别服务模块
    # ================================================================
    doc.add_heading('四、云端识别服务模块', level=1)
    add_table(doc,
        ['功能', '技术方案'],
        [
            ['主 OCR 引擎', 'PaddleOCR (PP-OCRv4) 自建服务，GPU 加速'],
            ['备用 OCR 引擎', '腾讯云 LicensePlateOCR API，主服务不可用时自动切换'],
            ['模糊匹配', 'SequenceMatcher 字符串相似度，阈值 0.70，匹配数据库注册车牌'],
            ['多帧去抖', '滑动窗口（最近 5 帧 ≥ 3 帧一致才确认）'],
            ['未匹配处理', '通知安保人工介入 + 抓拍图弹窗'],
        ],
        col_widths=[4, 13]
    )

    # ================================================================
    # 五、后端服务模块
    # ================================================================
    doc.add_heading('五、后端服务模块', level=1)

    doc.add_heading('5.1 框架与中间件', level=2)
    add_table(doc,
        ['组件', '技术方案'],
        [
            ['Web 框架', 'FastAPI (异步、自动 OpenAPI 文档、Pydantic 校验) 或 Flask'],
            ['ORM', 'SQLAlchemy'],
            ['数据库迁移', 'Alembic'],
            ['认证鉴权', 'JWT (flask-jwt-extended / fastapi-jwt-auth)'],
            ['接口限流', 'flask-limiter / slowapi'],
            ['跨域处理', 'flask-cors / fastapi-cors (白名单)'],
            ['日志记录', 'structlog + RotatingFileHandler (按日轮转，保留30天)'],
            ['异常处理', '全局 error_handler 统一 JSON 响应 {code, message, data}'],
            ['入参校验', 'Pydantic / marshmallow'],
        ],
        col_widths=[3.5, 13.5]
    )

    doc.add_heading('5.2 架构分层', level=2)
    add_table(doc,
        ['层级', '职责', '技术'],
        [
            ['路由层', '接收 HTTP/WS 请求，参数解析，响应封装', 'FastAPI Router / Flask Blueprint'],
            ['服务层', '业务逻辑编排（停车入场/出场、计费、推荐、审批）', '纯 Python 类，可单元测试'],
            ['数据层', '数据库操作封装', 'SQLAlchemy ORM Model'],
        ],
        col_widths=[2.5, 8, 6.5]
    )

    doc.add_heading('5.3 核心 API 清单', level=2)
    add_table(doc,
        ['模块', '接口', '方法', '说明'],
        [
            ['认证', '/api/auth/login', 'POST', '用户登录，返回 JWT token'],
            ['认证', '/api/auth/register', 'POST', '用户注册（待审批）'],
            ['认证', '/api/auth/refresh', 'POST', '刷新 access_token'],
            ['用户', '/api/user/info', 'GET', '获取个人信息'],
            ['用户', '/api/user/recommend', 'GET', '获取推荐车位'],
            ['用户', '/api/user/reserve', 'POST', '预约车位'],
            ['用户', '/api/user/records', 'GET', '我的停车记录'],
            ['用户', '/api/user/pay', 'POST', '支付停车费'],
            ['管理', '/api/admin/dashboard', 'GET', '管理端仪表盘数据'],
            ['管理', '/api/admin/users', 'GET', '用户列表'],
            ['管理', '/api/admin/approve', 'POST', '审批用户注册'],
            ['管理', '/api/admin/records', 'GET', '全场停车记录'],
            ['硬件', '/api/hardware/upload', 'POST', '边缘端上传识别结果/图像'],
            ['硬件', '/api/hardware/update_slot', 'POST', '传感器上报车位状态'],
            ['硬件', '/api/hardware/heartbeat', 'POST', '设备心跳上报'],
        ],
        col_widths=[1.5, 5, 1.5, 9]
    )

    # ================================================================
    # 六、数据库模块
    # ================================================================
    doc.add_heading('六、数据库模块', level=1)

    doc.add_heading('6.1 数据库选型', level=2)
    add_table(doc,
        ['类型', '选型', '用途'],
        [
            ['关系型数据库', 'PostgreSQL 16', '用户、车位、停车记录、计费规则、审计日志'],
            ['缓存', 'Redis 7', '车位实时状态、推荐结果缓存、JWT 黑名单'],
        ],
        col_widths=[3.5, 4, 9.5]
    )

    doc.add_heading('6.2 数据表清单', level=2)
    add_table(doc,
        ['表名', '用途'],
        [
            ['users', '用户账户'],
            ['parking_slots', '车位状态'],
            ['parking_records', '停车流水记录'],
            ['charging_records', '充电记录'],
            ['modification_requests', '用户信息变更审批'],
            ['parking_price_rules', '计费规则配置'],
            ['audit_logs', '操作审计日志'],
            ['hardware_devices', '硬件设备注册'],
        ],
        col_widths=[5, 12]
    )

    # ================================================================
    # 七、AI 预测与推荐模块
    # ================================================================
    doc.add_heading('七、AI 预测与推荐模块', level=1)
    add_table(doc,
        ['功能', '技术方案', '数据来源'],
        [
            ['车流预测', 'LSTM 时序预测 / Prophet', 'ParkingRecord 历史占用数据'],
            ['车位推荐', '协同过滤 (NearestNeighbors + Cosine Similarity)', '用户历史停车偏好'],
            ['新用户冷启动', '规则推荐：距离最近 + 充电桩优先 (绿牌)', '—'],
            ['模型持久化', 'joblib 序列化，启动加载，每日凌晨增量训练', '—'],
        ],
        col_widths=[3, 7.5, 6.5]
    )

    # ================================================================
    # 八、前端模块
    # ================================================================
    doc.add_heading('八、前端模块', level=1)

    doc.add_heading('8.1 技术栈', level=2)
    add_table(doc,
        ['技术', '选型'],
        [
            ['框架', 'Vue 3 (Composition API)'],
            ['构建工具', 'Vite'],
            ['类型系统', 'TypeScript'],
            ['UI 组件库', 'Element Plus'],
            ['可视化图表', 'ECharts 5'],
            ['状态管理', 'Pinia'],
            ['路由', 'Vue Router 4'],
            ['HTTP 请求', 'Axios + 拦截器 (JWT 自动附带)'],
            ['实时通信', 'socket.io-client'],
            ['移动端适配', 'postcss-px-to-viewport'],
        ],
        col_widths=[4, 13]
    )

    doc.add_heading('8.2 页面清单', level=2)
    add_table(doc,
        ['页面', '目标用户', '核心功能'],
        [
            ['登录/注册', '所有用户', '手机号+密码，JWT 认证'],
            ['用户首页', '车主', '车位地图、推荐车位、预约、充电管理'],
            ['停车记录', '车主', '历史停车、费用明细'],
            ['管理仪表盘', '管理员', 'KPI 卡片、实时占用率、今日营收'],
            ['用户/审批管理', '管理员', '用户审批、角色分配、车牌修改审批'],
            ['停车记录管理', '管理员', '全场记录查询、异常处理、无牌车登记'],
            ['数据报表', '管理员', '营收趋势、车位周转率、高峰时段分析'],
            ['实时大屏', '展厅/保安室', '全屏监控、实时车位 + 识别画面流'],
            ['设备监控', '运维', '设备在线状态、识别成功率、帧率'],
        ],
        col_widths=[3, 2.5, 11.5]
    )

    # ================================================================
    # 九、硬件通信模块
    # ================================================================
    doc.add_heading('九、硬件通信模块', level=1)

    doc.add_heading('9.1 通信方式', level=2)
    add_table(doc,
        ['模式', '技术方案', '适用场景'],
        [
            ['本地串口', 'PySerial，USB 连接 ESP32 ↔ PC', '单点部署 / 开发调试'],
            ['远程 MQTT', 'MQTT Broker (EMQX)，ESP32 WiFi 直连云端', '多设备分布式部署'],
        ],
        col_widths=[2.5, 7, 7.5]
    )

    doc.add_heading('9.2 应用层协议', level=2)
    doc.add_paragraph('协议格式：CMD|SEQ|PAYLOAD|CHECKSUM')
    add_table(doc,
        ['指令', '方向', '说明'],
        [
            ['ENTRY', '设备 → 云端', '车辆入场上报（车牌 + 时间戳）'],
            ['EXIT', '设备 → 云端', '车辆出场上报'],
            ['ACK_ENTRY / ACK_EXIT', '云端 → 设备', '入场/出场确认（含费用）'],
            ['PING / PONG', '双向', '心跳保活'],
            ['SLOT', '设备 → 云端', '车位状态上报（地磁传感器）'],
            ['TIME', '云端 → 设备', '时间同步'],
            ['ERR', '双向', '错误上报'],
        ],
        col_widths=[4, 3, 10]
    )

    doc.add_heading('9.3 容错机制', level=2)
    add_table(doc,
        ['机制', '技术方案'],
        [
            ['消息校验', 'CRC-16 校验码'],
            ['超时重传', '3 秒无 ACK 则重发，最多 3 次'],
            ['断线重连', '指数退避 (1s→2s→4s→8s→max 30s)'],
            ['离线缓存', 'ESP32 Flash 暂存，恢复后补传'],
        ],
        col_widths=[4, 13]
    )

    # ================================================================
    # 十、计费引擎模块
    # ================================================================
    doc.add_heading('十、计费引擎模块', level=1)
    add_table(doc,
        ['计费项', '规则'],
        [
            ['免费', '入场 30 分钟内免费'],
            ['白天 (8:00-20:00)', '¥5/小时，封顶 ¥50/天'],
            ['夜间 (20:00-8:00)', '¥3/小时，封顶 ¥20/夜'],
            ['新能源绿牌', '享 8 折优惠'],
            ['计费规则存储', 'parking_price_rules 表，管理员可后台修改即时生效'],
        ],
        col_widths=[5, 12]
    )

    # ================================================================
    # 十一、实时推送模块
    # ================================================================
    doc.add_heading('十一、实时推送模块 (WebSocket)', level=1)
    add_table(doc,
        ['推送事件', '内容', '触发时机'],
        [
            ['slot_update', '车位号、分区、状态、车牌', '车位状态变化'],
            ['entry_alert / exit_alert', '车牌、车位、费用、时间', '车辆入场/出场'],
            ['dashboard_stats', '总车位数、空闲、占用、今日营收', '每 5 秒'],
            ['recognition_feed', '抓拍图 Base64、车牌、置信度', '每次识别完成'],
            ['system_alert', '告警级别、标题、描述', '满位/设备离线/异常'],
        ],
        col_widths=[4, 7, 6]
    )

    # ================================================================
    # 十二、部署与运维模块
    # ================================================================
    doc.add_heading('十二、部署与运维模块', level=1)

    doc.add_heading('12.1 部署方案', level=2)
    add_table(doc,
        ['组件', '技术方案'],
        [
            ['容器化', 'Docker + Docker Compose'],
            ['Web 服务器', 'Gunicorn (4 workers)'],
            ['反向代理 + HTTPS', 'Nginx + Let\'s Encrypt (certbot 自动续签)'],
            ['数据库', 'PostgreSQL 16 + Redis 7'],
            ['CI/CD', 'GitHub Actions: pytest → ruff → mypy → docker build'],
        ],
        col_widths=[4, 13]
    )

    doc.add_heading('12.2 监控体系', level=2)
    add_table(doc,
        ['工具', '用途'],
        [
            ['Prometheus + Grafana', '请求量、响应时间、车位占用率、告警规则'],
            ['Sentry', 'Python 异常 + 前端 JS 错误追踪'],
            ['Uptime Kuma', 'HTTP 服务存活检测 + SSL 到期提醒'],
        ],
        col_widths=[5, 12]
    )

    doc.add_heading('12.3 数据备份', level=2)
    doc.add_paragraph(
        'PostgreSQL 每日凌晨 pg_dump 全量备份，保留 30 天，同步至对象存储。每月一次异地冷备。'
    )

    # ================================================================
    # 十三、测试方案
    # ================================================================
    doc.add_heading('十三、测试方案', level=1)
    add_table(doc,
        ['测试层次', '工具', '覆盖目标'],
        [
            ['单元测试', 'pytest + pytest-cov', 'Service 层业务逻辑，覆盖率 > 80%'],
            ['集成测试', 'pytest + TestClient', 'API 接口 + 数据库完整链路'],
            ['端到端测试', 'Playwright', '关键用户流程（登录→预约→支付）'],
        ],
        col_widths=[3.5, 5, 8.5]
    )

    # ================================================================
    # 十四、项目里程碑
    # ================================================================
    doc.add_heading('十四、项目里程碑', level=1)
    add_table(doc,
        ['阶段', '周期', '交付物', '验收标准'],
        [
            ['Phase 1: 基础架构', '第1-2周', 'JWT 认证、项目重构、配置管理、日志', 'API 文档完整、认证通过'],
            ['Phase 2: 数据与测试', '第3-4周', 'PostgreSQL 迁移、Alembic、Service 层测试', '覆盖率 80%+、集成测试通过'],
            ['Phase 3: AI 引擎', '第5周', '模型持久化、真实数据训练', '预测 MAE < 15%'],
            ['Phase 4: 前端', '第6-7周', 'Vue 3 重构、WebSocket 推送、报表页面', '页面加载 < 2s'],
            ['Phase 5: 硬件部署', '第8周', '串口容错、MQTT、设备监控 API', '断线自动重连正常'],
            ['Phase 6: 部署上线', '第9周', 'Docker 化、CI/CD、监控面板', '可用性 > 99.5%'],
        ],
        col_widths=[3.5, 2, 5.5, 6]
    )

    # ===== 保存 =====
    output_path = os.path.join(os.path.dirname(__file__), '智能停车场管理系统_项目设计书.docx')
    doc.save(output_path)
    print(f'文档已生成: {output_path}')

if __name__ == '__main__':
    generate_doc()
